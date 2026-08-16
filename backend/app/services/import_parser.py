import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

FIELD_ALIASES = {
    "sale_date": {"sale_date", "date", "日期", "销售日期"},
    "snapshot_date": {"snapshot_date", "date", "日期", "库存日期", "快照日期"},
    "sku": {"sku", "asin", "商品sku", "商品编码"},
    "product_name": {"product_name", "name", "商品名称", "产品名称"},
    "category": {"category", "品类", "类目"},
    "site": {"site", "站点", "国家站"},
    "platform": {"platform", "平台"},
    "units_sold": {"units_sold", "quantity", "销量", "销售数量"},
    "revenue": {"revenue", "sales", "销售额", "收入"},
    "cost": {"cost", "成本", "销售成本"},
    "refund_units": {"refund_units", "refunds", "退款数量", "退货数量"},
    "on_hand": {"on_hand", "库存", "在库", "现有库存"},
    "reserved": {"reserved", "预留", "锁定库存"},
    "inbound": {"inbound", "在途", "入库中"},
}

REQUIRED_FIELDS = {
    "sales": [
        "sale_date",
        "sku",
        "product_name",
        "category",
        "site",
        "platform",
        "units_sold",
        "revenue",
        "cost",
        "refund_units",
    ],
    "inventory": [
        "snapshot_date",
        "sku",
        "product_name",
        "category",
        "site",
        "platform",
        "on_hand",
        "reserved",
        "inbound",
    ],
}


@dataclass
class RowIssue:
    row_number: int
    field_name: str | None
    error_code: str
    message: str
    raw_row: dict[str, Any]


@dataclass
class ParsedImport:
    headers: list[str]
    mapping: dict[str, str]
    valid_rows: list[dict[str, Any]]
    issues: list[RowIssue]
    total_rows: int


def read_tabular_file(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            headers = [str(header).strip() for header in (reader.fieldnames or [])]
            rows = [{str(key).strip(): value for key, value in row.items()} for row in reader]
            return headers, rows

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    first_row = next(values, None)
    if first_row is None:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in first_row]
    rows = []
    for values_row in values:
        rows.append({header: value for header, value in zip(headers, values_row, strict=False)})
    workbook.close()
    return headers, rows


def detect_mapping(headers: list[str], kind: str) -> dict[str, str]:
    normalized_headers = {header.strip().lower(): header for header in headers}
    mapping: dict[str, str] = {}
    for canonical in REQUIRED_FIELDS[kind]:
        for alias in FIELD_ALIASES[canonical]:
            if alias.lower() in normalized_headers:
                mapping[canonical] = normalized_headers[alias.lower()]
                break
    return mapping


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m", "%Y/%m"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("日期格式应为 YYYY-MM-DD")


def _parse_int(value: Any, field: str) -> int:
    try:
        parsed = int(str(value).strip())
    except (TypeError, ValueError) as error:
        raise ValueError(f"{field} 必须是整数") from error
    if parsed < 0:
        raise ValueError(f"{field} 不能小于 0")
    return parsed


def _parse_decimal(value: Any, field: str) -> Decimal:
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError) as error:
        raise ValueError(f"{field} 必须是数字") from error
    if parsed < 0:
        raise ValueError(f"{field} 不能小于 0")
    return parsed.quantize(Decimal("0.01"))


def validate_import_rows(
    headers: list[str],
    rows: list[dict[str, Any]],
    kind: str,
    mapping: dict[str, str] | None = None,
) -> ParsedImport:
    effective_mapping = mapping or detect_mapping(headers, kind)
    issues: list[RowIssue] = []
    valid_rows: list[dict[str, Any]] = []
    missing_fields = [field for field in REQUIRED_FIELDS[kind] if field not in effective_mapping]
    if missing_fields:
        for field in missing_fields:
            issues.append(RowIssue(1, field, "missing_column", f"缺少必填字段：{field}", {}))
        return ParsedImport(headers, effective_mapping, [], issues, len(rows))

    seen: set[tuple[str, str, str, str]] = set()
    for row_number, row in enumerate(rows, start=2):
        normalized: dict[str, Any] = {}
        row_issues: list[RowIssue] = []
        for field in REQUIRED_FIELDS[kind]:
            source_field = effective_mapping[field]
            value = row.get(source_field)
            if value is None or str(value).strip() == "":
                row_issues.append(RowIssue(row_number, field, "required", f"{field} 不能为空", row))
                continue
            try:
                if field in {"sale_date", "snapshot_date"}:
                    normalized[field] = _parse_date(value)
                elif field in {"units_sold", "refund_units", "on_hand", "reserved", "inbound"}:
                    normalized[field] = _parse_int(value, field)
                elif field in {"revenue", "cost"}:
                    normalized[field] = _parse_decimal(value, field)
                else:
                    normalized[field] = str(value).strip()
            except ValueError as error:
                row_issues.append(RowIssue(row_number, field, "invalid_format", str(error), row))

        if not row_issues:
            date_field = "sale_date" if kind == "sales" else "snapshot_date"
            duplicate_key = (
                normalized[date_field].isoformat(),
                normalized["sku"],
                normalized["site"],
                normalized["platform"],
            )
            if duplicate_key in seen:
                row_issues.append(
                    RowIssue(row_number, None, "duplicate_row", "文件内存在重复记录", row)
                )
            else:
                seen.add(duplicate_key)

        if row_issues:
            issues.extend(row_issues)
        else:
            valid_rows.append({"source_row_number": row_number, **normalized})

    return ParsedImport(headers, effective_mapping, valid_rows, issues, len(rows))


def parse_and_validate(
    path: Path, kind: str, mapping: dict[str, str] | None = None
) -> ParsedImport:
    headers, rows = read_tabular_file(path)
    return validate_import_rows(headers, rows, kind, mapping)
